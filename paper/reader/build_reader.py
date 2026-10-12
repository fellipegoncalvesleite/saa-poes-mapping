"""Build the personal Portuguese reader version (not for publication) as DOCX and PDF.

Run from repository root, after scripts/build_claims_register.py and scripts/make_paper_figures.py:

    python paper/reader/build_reader.py

Requires python-docx and typst (paper/reader/requirements-reader.txt). The text lives in
paper/reader/conteudo.py; every number is a ``{key}`` placeholder resolved from
paper/generated/claims_register.json, so the reader cannot drift from the manuscript. A missing key
stops the build. Writes paper/reader/build/SAA_POES_leitura_pessoal_PT.{docx,pdf}.
"""
from __future__ import annotations

from pathlib import Path
import json
import re
import sys

import typst
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
import conteudo  # noqa: E402

OUT = HERE / "build"
NAME = "SAA_POES_leitura_pessoal_PT"
NUM = re.compile(r"^[−+-]?[\d \u202f\u2009\u00a0]+(\.\d+)?$")
IDENTIFIERS = {"ev_index", "ev_utc", "ev_sha"}  # printed verbatim, never as quantities


def pt(formatted: str) -> str:
    """Register display value → Brazilian convention (1.234.567,8); non-numeric values unchanged."""
    if not NUM.match(formatted):
        return formatted
    sign = formatted[0] if formatted[0] in "−+-" else ""
    body = re.sub(r"[ \u202f\u2009\u00a0]", "", formatted[len(sign):])
    whole, _, frac = body.partition(".")
    groups = []
    while len(whole) > 3:
        groups.insert(0, whole[-3:])
        whole = whole[:-3]
    groups.insert(0, whole)
    return sign + ".".join(groups) + ("," + frac if frac else "")


def values() -> dict:
    reg = json.loads((ROOT / "paper/generated/claims_register.json").read_text())
    v = {c["key"]: c["formatted"] if c["key"] in IDENTIFIERS else pt(c["formatted"]) for c in reg}
    v.update({k: pt(x) if isinstance(x, str) else x for k, x in conteudo.EXTRA.items()})
    return v


def fill(text: str, v: dict) -> str:
    def sub(m):
        if m.group(1) not in v:
            raise KeyError(f"reader placeholder {{{m.group(1)}}} is not in the claims register")
        return v[m.group(1)]
    return re.sub(r"\{([A-Za-z0-9_]+)\}", sub, text)


# ----------------------------------------------------------------------------------------- DOCX
def runs(par, text: str):
    for i, piece in enumerate(re.split(r"\*\*", text)):
        r = par.add_run(piece)
        r.bold = bool(i % 2)
    return par


def shade(cell, hex_fill):
    tc = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_fill)
    tc.append(sh)


def docx(blocks, v):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
        hp = s.header.paragraphs[0]
        hp.text = "Versão pessoal de leitura — não é para publicação"
        hp.runs[0].font.size = Pt(8)
        hp.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    for b in blocks:
        kind = b[0]
        if kind == "title":
            p = d.add_paragraph()
            r = p.add_run(fill(b[1], v))
            r.bold, r.font.size = True, Pt(20)
        elif kind == "subtitle":
            p = d.add_paragraph()
            r = p.add_run(fill(b[1], v))
            r.font.size, r.italic = Pt(12), True
        elif kind in ("h1", "h2", "h3"):
            d.add_heading(fill(b[1], v), level={"h1": 1, "h2": 2, "h3": 3}[kind])
        elif kind == "p":
            runs(d.add_paragraph(), fill(b[1], v))
        elif kind in ("ul", "ol"):
            for item in b[1]:
                runs(d.add_paragraph(style="List Bullet" if kind == "ul" else "List Number"), fill(item, v))
        elif kind == "box":
            t = d.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            c = t.rows[0].cells[0]
            shade(c, "FFF4E5")
            c.paragraphs[0].text = ""
            runs(c.paragraphs[0], fill(b[1], v))
            d.add_paragraph()
        elif kind == "table":
            head, rows = b[1], b[2]
            t = d.add_table(rows=1 + len(rows), cols=len(head))
            t.style = "Table Grid"
            for j, h in enumerate(head):
                c = t.rows[0].cells[j]
                c.paragraphs[0].text = ""
                rr = c.paragraphs[0].add_run(fill(h, v))
                rr.bold, rr.font.size = True, Pt(9.5)
                shade(c, "E8EEF4")
            for i, row in enumerate(rows, start=1):
                for j, cell in enumerate(row):
                    c = t.rows[i].cells[j]
                    c.paragraphs[0].text = ""
                    for k, piece in enumerate(re.split(r"\*\*", fill(cell, v))):
                        rr = c.paragraphs[0].add_run(piece)
                        rr.bold, rr.font.size = bool(k % 2), Pt(9.5)
            d.add_paragraph()
        elif kind == "figure":
            d.add_picture(str(ROOT / b[1]), width=Cm(16))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = d.add_paragraph()
            r = cap.add_run(fill(b[2], v))
            r.italic, r.font.size = True, Pt(9.5)
        elif kind == "eq":
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(fill(b[2], v))
            r.font.name = "Cambria Math"
        elif kind == "pagebreak":
            d.add_page_break()
        else:
            raise ValueError(kind)
    OUT.mkdir(exist_ok=True)
    d.core_properties.author = "Fellipe Gonçalves Leite"
    d.core_properties.title = "SAA POES — versão pessoal de leitura (PT)"
    d.save(OUT / f"{NAME}.docx")


# ----------------------------------------------------------------------------------------- PDF (Typst)
def esc(text: str) -> str:
    out = re.sub(r'([\\#$@<>\[\]`_~*"])', r"\\\1", text).replace("//", "\\/\\/")
    return re.sub(r"\\\*\\\*(.+?)\\\*\\\*", r"*\1*", out)


def typ(blocks, v) -> str:
    lines = ['#set document(title: "SAA POES — versão pessoal de leitura", author: "Fellipe Gonçalves Leite")',
             '#set page(paper: "a4", margin: (x: 2.2cm, y: 2.2cm), numbering: "1",',
             '  header: align(right, text(7.5pt, fill: luma(120))[Versão pessoal de leitura — não é para publicação]))',
             '#set text(lang: "pt", region: "br", size: 10.5pt)',
             '#set par(justify: true, leading: 0.62em)',
             '#show heading.where(level: 1): it => block(above: 1.4em, below: 0.7em, text(14pt, it))',
             '#show heading.where(level: 2): it => block(above: 1.1em, below: 0.5em, text(11.5pt, it))',
             '#show heading.where(level: 3): it => block(above: 0.9em, below: 0.4em, text(10.5pt, it))',
             '#set table(stroke: 0.4pt + luma(170), inset: 4pt)']
    for b in blocks:
        kind = b[0]
        if kind == "title":
            lines.append(f"#text(19pt, weight: \"bold\")[{esc(fill(b[1], v))}]\n")
        elif kind == "subtitle":
            lines.append(f"#text(11.5pt, style: \"italic\")[{esc(fill(b[1], v))}]\n")
        elif kind in ("h1", "h2", "h3"):
            lines.append("=" * int(kind[1]) + " " + esc(fill(b[1], v)) + "\n")
        elif kind == "p":
            lines.append(esc(fill(b[1], v)) + "\n")
        elif kind in ("ul", "ol"):
            mark = "-" if kind == "ul" else "+"
            lines.extend(f"{mark} {esc(fill(i, v))}" for i in b[1])
            lines.append("")
        elif kind == "box":
            lines.append(f"#block(fill: rgb(\"#fff4e5\"), stroke: 0.6pt + rgb(\"#b45309\"), inset: 8pt, radius: 3pt, "
                         f"width: 100%)[{esc(fill(b[1], v))}]\n")
        elif kind == "table":
            head, rows = b[1], b[2]
            cols = b[3] if len(b) > 3 else ["1fr"] * len(head)
            cells = ["table.header(" + ", ".join(f"[*{esc(fill(h, v))}*]" for h in head) + ")"]
            for row in rows:
                cells.extend(f"[{esc(fill(c, v))}]" for c in row)
            lines.append(f"#block(breakable: true, text(9pt, table(columns: ({', '.join(cols)}), "
                         f"fill: (_, y) => if y == 0 {{ rgb(\"#e8eef4\") }}, {', '.join(cells)})))\n")
        elif kind == "figure":
            rel = Path("../../..") / b[1]
            cap = re.sub(r"^Figura \d+\.\s*", "", fill(b[2], v))  # Typst numbers figures itself
            lines.append(f"#figure(image(\"{rel.as_posix()}\", width: 100%), caption: [{esc(cap)}])\n")
        elif kind == "eq":
            lines.append(f"$ {fill(b[1], v)} $\n")
        elif kind == "pagebreak":
            lines.append("#pagebreak()\n")
    return "\n".join(lines)


def main() -> None:
    v = values()
    blocks = conteudo.BLOCKS
    docx(blocks, v)
    OUT.mkdir(exist_ok=True)
    src = OUT / f"{NAME}.typ"
    src.write_text(typ(blocks, v))
    pdf, warnings = typst.compile_with_warnings(str(src), root=str(ROOT), ignore_system_fonts=True,
                                               timestamp=1790985600)  # 2026-10-03T00:00Z
    for w in warnings:
        print("warning:", w, file=sys.stderr)
    (OUT / f"{NAME}.pdf").write_bytes(pdf)
    print("built", OUT / f"{NAME}.pdf", OUT / f"{NAME}.docx")


if __name__ == "__main__":
    main()
